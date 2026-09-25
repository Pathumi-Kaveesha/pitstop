# Copyright (c) 2026 WSO2 LLC. (https://www.wso2.com).
#
# WSO2 LLC. licenses this file to you under the Apache License,
# Version 2.0 (the "License"); you may not use this file except
# in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

"""Extracts text per format, joins it, and splits it into chunks tagged
with the page/slide/section/sheet they came from.

One "piece" per format: PDF -> page, PPTX -> slide, DOCX -> heading-delimited
section (no page numbers in .docx), XLSX -> sheet tab.
"""

import io
import zipfile
from dataclasses import dataclass

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from config import (
    MAX_CHUNK_OVERLAP,
    MAX_CHUNK_SIZE,
    MAX_OOXML_UNCOMPRESSED_BYTES,
    MIN_CHUNK_LENGTH,
)

UNIT_LABELS = {
    "pdf": "Page",
    "pptx": "Slide",
    "docx": "Section",
    "xlsx": "Sheet",
}

@dataclass
class Chunk:
    text: str
    page: int


def _extract_pdf(file_bytes: bytes) -> list[str]:
    """One string per page."""
    reader = PdfReader(io.BytesIO(file_bytes))
    return [page.extract_text() or "" for page in reader.pages]


def _reject_oversized_ooxml(file_bytes: bytes) -> None:
    """Office files are zip archives - a small one can expand to something
    huge once parsed. The entries declare their own uncompressed sizes, so
    this reads them without decompressing anything."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            total = sum(entry.file_size for entry in archive.infolist())
    except zipfile.BadZipFile as error:
        raise ValueError("This file isn't a readable Office document.") from error

    if total > MAX_OOXML_UNCOMPRESSED_BYTES:
        raise ValueError(
            f"This document expands to {total / 1_048_576:.0f} MB, over the "
            f"{MAX_OOXML_UNCOMPRESSED_BYTES / 1_048_576:.0f} MB limit Smart Search can read."
        )


def _extract_pptx(file_bytes: bytes) -> list[str]:
    """One string per slide, including speaker notes."""
    _reject_oversized_ooxml(file_bytes)
    presentation = Presentation(io.BytesIO(file_bytes))

    slides = []
    for slide in presentation.slides:
        parts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(notes.text)

        slides.append("\n".join(parts))
    return slides


def _iter_docx_blocks(document):
    """Yields paragraphs and tables in document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _extract_docx(file_bytes: bytes) -> tuple[list[str], list[str | None]]:
    """One string per heading-delimited section, plus each section's heading."""
    _reject_oversized_ooxml(file_bytes)
    document = Document(io.BytesIO(file_bytes))

    sections: list[str] = []
    headings: list[str | None] = []
    current: list[str] = []
    current_heading: str | None = None
    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            rows = [
                ", ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                for row in block.rows
            ]
            current.extend(row for row in rows if row)
            continue

        text = block.text.strip()
        if not text:
            continue

        if block.style.name.startswith("Heading"):
            if current:
                sections.append("\n".join(current))
                headings.append(current_heading)
                current = []
            current_heading = text
        current.append(text)

    if current:
        sections.append("\n".join(current))
        headings.append(current_heading)
    return sections, headings


def _extract_xlsx(file_bytes: bytes) -> list[str]:
    """One string per sheet, rows written as "Header: value" pairs."""
    _reject_oversized_ooxml(file_bytes)
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)

    sheets = []
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)

        headers = next(rows, None)
        if headers is None:
            sheets.append("")
            continue
        header_names = [str(h).strip() if h is not None else "" for h in headers]

        lines = [f"Sheet: {worksheet.title}"]
        for row in rows:
            pairs = [
                f"{name}: {value}"
                for name, value in zip(header_names, row)
                if name and value is not None and str(value).strip()
            ]
            if pairs:
                lines.append(", ".join(pairs))
        sheets.append("\n".join(lines))

    workbook.close()
    return sheets


_EXTRACTORS = {
    "pdf": _extract_pdf,
    "pptx": _extract_pptx,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
}


def extract_units(file_bytes: bytes, extension: str) -> tuple[list[str], list[str | None]]:
    """One string per piece of the document, plus each piece's heading (docx only)."""
    result = _EXTRACTORS[extension](file_bytes)
    if extension == "docx":
        return result
    return result, [None] * len(result)


def chunk_document(file_bytes: bytes, extension: str) -> tuple[list[Chunk], list[str | None]]:
    """Extract -> join -> split -> tag with page -> drop short chunks.
    Also returns each unit's heading, for building deep links."""
    unit_texts, unit_headings = extract_units(file_bytes, extension)

    full_text = ""
    unit_start_offsets: list[int] = []
    for unit_text in unit_texts:
        unit_start_offsets.append(len(full_text))
        full_text += unit_text + "\n"

    if not full_text.strip():
        raise ValueError("No readable text found in this document")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=MAX_CHUNK_SIZE,
        chunk_overlap=MAX_CHUNK_OVERLAP,
    )
    raw_chunks = splitter.split_text(full_text)

    # Chunks come back in document order, so each one starts at or after the
    # previous one. Searching from there stops repeated text (a boilerplate
    # slide, a repeated header) from matching its first occurrence and
    # reporting the wrong page. The cursor advances for skipped chunks too.
    chunks: list[Chunk] = []
    search_from = 0
    for text in raw_chunks:
        offset = full_text.find(text, search_from)
        if offset == -1:
            offset = full_text.find(text)
        if offset != -1:
            search_from = offset + 1

        if len(text.strip()) < MIN_CHUNK_LENGTH:
            continue
        chunks.append(Chunk(text=text, page=find_page_number(offset, unit_start_offsets)))
    return chunks, unit_headings


def find_page_number(offset: int, unit_start_offsets: list[int]) -> int:
    """Which piece a chunk came from, by its offset in the joined text."""
    if offset == -1:
        return 1

    page = 1
    for i, start in enumerate(unit_start_offsets):
        if start <= offset:
            page = i + 1
        else:
            break
    return page
